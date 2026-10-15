from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image,ImageDraw,ImageFont
R=Path(__file__).resolve().parents[2]
font=str(R/'Documentation/fonts/DejaVuSans.ttf')
def diagram(name,workflow=False):
 im=Image.new('RGB',(1500,620),'white');d=ImageDraw.Draw(im);f=ImageFont.truetype(font,25);small=ImageFont.truetype(font,21)
 def box(x,y,w,h,lines):
  d.rounded_rectangle((x,y,x+w,y+h),12,fill='#edf4f7',outline='#536a78',width=3)
  for i,t in enumerate(lines):d.text((x+w/2,y+19+i*32),t,font=f,fill='#102d3b',anchor='mt')
 def arrow(a,b):
  d.line([a,b],fill='#536a78',width=4);x,y=b
  if b[0]==a[0]:d.polygon([(x,y),(x-9,y-16),(x+9,y-16)],fill='#536a78')
  else:d.polygon([(x,y),(x-16,y-9),(x-16,y+9)],fill='#536a78')
 if not workflow:
  box(480,20,540,85,['Browser AI Assessment page']);box(480,155,540,85,['FastAPI and background job']);arrow((750,105),(750,155))
  box(30,335,430,115,['Evidence retrieval','BM25 and token hashes']);box(540,335,430,115,['Ollama local inference','Qwen2.5 3B Instruct']);box(1040,335,430,115,['PostgreSQL and pgvector','ai_assessments table'])
  d.line([(750,240),(750,280),(245,280),(245,335)],fill='#536a78',width=4);arrow((750,240),(750,335));d.line([(750,280),(1255,280),(1255,335)],fill='#536a78',width=4)
  d.text((750,510),'Evidence snapshot → schema checked draft → saved record → human review',font=small,fill='#102d3b',anchor='mt')
 else:
  box(40,40,420,100,['New system description','Input sources SYS-xx']);box(540,40,420,100,['Retrieve relevant guidance','Top 4 reference passages']);box(1040,40,420,100,['Generate structured JSON','Assets and threat scenarios']);arrow((460,90),(540,90));arrow((960,90),(1040,90))
  box(1040,305,420,100,['Check schema and citations','Reject invalid model output']);box(540,305,420,100,['Save to PostgreSQL','Status Draft or Failed']);box(40,305,420,100,['Review and export','No automatic risk approval']);arrow((1250,140),(1250,305));d.line([(1040,355),(960,355)],fill='#536a78',width=4);d.line([(540,355),(460,355)],fill='#536a78',width=4)
  d.text((750,510),'Failed inference is reported explicitly. No predefined assessment is substituted.',font=small,fill='#102d3b',anchor='mt')
 im.save(R/'Documentation/figures'/name)
diagram('architecture.png');diagram('workflow.png',True)
def new(template):
 d=Document(template)
 for el in list(d.element.body):
  if el.tag!=qn('w:sectPr'):d.element.body.remove(el)
 for name in ['Normal','normal','Heading 1','Heading 2','Title']:
  try:
   st=d.styles[name];st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0)
   st.font.size=Pt(11 if name.lower()=='normal' else 14 if name.startswith('Heading') else 19)
  except KeyError:pass
 normal=d.styles['Normal'];normal.paragraph_format.space_after=Pt(7);normal.paragraph_format.line_spacing=1.08
 for sec in d.sections:
  sec.header_distance=Inches(.3);sec.footer_distance=Inches(.3)
  p=sec.footer.paragraphs[0];p.clear();p.alignment=2;r=p.add_run();fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');r._r.addnext(fld)
 for style in d.styles:
  if style.element.pPr is not None:
   for border in list(style.element.pPr.findall(qn('w:pBdr'))):style.element.pPr.remove(border)
 return d
def para(d,t,bold=False):
 p=d.add_paragraph();r=p.add_run(t);r.bold=bold;return p
def h(d,t):return d.add_heading(t,1)
def table(d,rows,widths=None):
 t=d.add_table(rows=0,cols=len(rows[0]));t.autofit=False
 if widths:
  for c,w in zip(t.columns,widths):c.width=Inches(w)
 for i,row in enumerate(rows):
  cells=t.add_row().cells
  for j,text in enumerate(row):
   cells[j].text=str(text)
   if widths:cells[j].width=Inches(widths[j])
   for p in cells[j].paragraphs:
    p.paragraph_format.space_after=Pt(5);p.paragraph_format.space_before=Pt(5)
    for r in p.runs:r.font.size=Pt(9.5);r.bold=i==0
   pr=cells[j]._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');b.append(e)
   pr.append(b)
   if i==0:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'DFEAF2');pr.append(sh)
  if i==0:flag=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(flag)
  no=OxmlElement('w:cantSplit');t.rows[-1]._tr.get_or_add_trPr().append(no)
 d.add_paragraph();return t
def brand(d,kind):
 for text in ['TATA TECHNOLOGIES LTD.','TECH PULSE FY-26 | AI/ML Capstone Project',kind]:
  p=d.add_paragraph(text,style='Title' if text==kind else 'Normal');p.alignment=1
  if text!=kind:
   for r in p.runs:r.bold=True;r.font.size=Pt(12)
def pic(d,path,width=6.3):d.add_picture(str(path),width=Inches(width))
def page(d):d.add_page_break()
D=new(str(R/'Documentation/templates/PCCOE_AIML_Technical_Report_Template.docx'));brand(D,'TECHNICAL REPORT')
table(D,[['Field','Details'],['Institute / University','PCCOE'],['Programme / Batch','____________________________'],['Student Name','____________________________'],['PRN','____________________________'],['Case Study ID (CS1–CS5)','CS3'],['Project Title','TARA Studio Automotive Cybersecurity Assessment Assistant']], [1.8,4.8])
h(D,'1. Problem Statement and Intended Users')
para(D,'Automotive cybersecurity engineers must connect system architecture, protected assets, threat scenarios, controls and verification evidence. Manual preparation can leave inconsistent assumptions and fragmented traceability. TARA Studio supports this work with source-grounded draft assessments and an explicit human-review workflow.')
para(D,'The v1.1 AI Assessment page accepts a new system description, retrieves supporting guidance, requests structured output from a local Ollama model and saves the draft to PostgreSQL when the provided AI deployment is used. Intended users are cybersecurity engineers, system architects, student project teams and faculty evaluators. The application does not certify a vehicle or replace expert risk assessment.')
h(D,'2. Input Data / Knowledge Base')
table(D,[['Item','Description'],['Sources','Original synthetic OTA architectures, threat handbook, pattern library and user-supplied system descriptions. No proprietary OEM data or standards clauses.'],['Size and format','12 handbook passages; 5 baseline patterns; 2 architecture JSON versions; 18 baseline evaluation questions; 2 fresh text descriptions for Bluetooth access and EV charging.'],['Preprocessing / EDA done','Text extraction, source locators, identifier validation and deterministic tokenization. No statistical model training or dataset EDA is claimed.'],['Chunking strategy','Knowledge uploads: up to 1,200 characters with 100-character overlap. AI input description: non-overlapping 1,000-character source chunks, maximum 8,000 characters.']], [1.6,5])
page(D);h(D,'3. Solution Architecture')
para(D,'FastAPI serves the browser interface and coordinates a single-worker background generation job. The new assessment route does not call the predefined threat-pattern engine. The offline workspace remains available separately for comparison and basic workflow demonstrations.')
table(D,[['Component','Choice name version and hosting'],['Pre-trained model','Qwen2.5 3B Instruct via local Ollama, tag qwen2.5:3b-instruct. Model weights must be downloaded separately. Exact digest is captured from the installed model when generation runs.'],['Embedding model','No pretrained embedding model in this release. Deterministic signed token-hash vectors, 384 dimensions, with BM25 lexical ranking.'],['Vector store (FAISS / ChromaDB)','PostgreSQL with pgvector in the supplied AI deployment. FAISS and ChromaDB are not used. SQLite remains the native fallback and is labelled in the interface.'],['RAG pipeline / AI-agent design','Retrieve up to 4 reference passages, add source-addressable system input, generate constrained JSON, validate references and save a draft. No autonomous tool-executing agent.'],['Service layer and UI','Python; FastAPI 0.115.12; Uvicorn 0.34.2; Pydantic 2.11.4; local HTML/CSS/JavaScript.'],['Tools / deterministic rules','pypdf 5.5.0; psycopg 3.2.9; JSON schema, citation and asset-link checks; review revision checks; educational risk matrix in the baseline only.']], [1.6,5])
pic(D,R/'Documentation/figures/architecture.png');para(D,'Architecture of the v1.1 generative assessment route. Local model and database services are provisioned by compose.ai.yaml.')
page(D);h(D,'4. Implementation and Configuration')
for t in [
'1. Ingestion: the user enters a title and a new automotive system description or loads a text file. The request accepts 100 to 8,000 characters and requests 1 to 6 threats; starting with 1 threat reduces initial CPU inference time.',
'2. Retrieval: the backend snapshots input as SYS-xx evidence and retrieves up to four relevant knowledge passages. The lexical ranker combines BM25 with a 0.25-weight token-hash cosine term. PostgreSQL mode retrieves up to 30 vector candidates before re-ranking.',
'3. Generation: Ollama /api/generate receives the description, evidence, system prompt and the Pydantic-generated JSON schema. Configuration is temperature 0, num_ctx 8192, num_predict 4096 and a 600-second default timeout. The model drafts assets, threat and damage scenarios, hypothesized attack paths, controls, assumptions, mitigations, requirements and verification proposals.',
'4. Response handling: the application checks schema validity, unique identifiers, asset references, source membership and at least one submitted-system citation per asset and threat. Invalid output becomes Failed; there is no canned success fallback. Citation membership does not establish semantic correctness.',
'5. Persistence: queued jobs, progress, model metadata, input, evidence, output and review state are written to ai_assessments. PostgreSQL uses parameterized SQL and conflict-safe record updates. The page polls the job and displays the saved result. JSON export preserves the evidence snapshot.',
'6. Human review: generated records begin as Draft and risk remains Unassessed. A reviewer can record Reviewed or Rejected with a rationale. This is not formal security acceptance. Previously stored offline architecture findings are not overwritten.',
'Deployment: compose.ai.yaml connects app, PostgreSQL/pgvector and Ollama on a local container network. Only the app is published on localhost port 8000. Initial package, image and model downloads need internet; subsequent inference uses local services. Model and image digests must be recorded on the demonstration machine.'
]:para(D,t)
pic(D,R/'Documentation/figures/workflow.png')
page(D);h(D,'5. Evaluation Steps and Results')
para(D,'The recorded results below come from earlier offline-baseline and software-fixture runs. Further test validation was not performed for the final update. No live Qwen generation quality, PostgreSQL deployment result or student review is represented as completed. The synthetic questions were authored with the reference handbook and are not an independent automotive benchmark.')
table(D,[['#','Test question / scenario','Expected answer / behaviour','Result Pass Fail notes'],['1','Prevent firmware rollback','Retrieve KB-02','Pass in recorded offline baseline'],['2','Correctly signed but obsolete release','Retrieve KB-02','Fail: KB-11 ranked first; expected source absent from top 3'],['3','Who may authorize software packages','Retrieve KB-03','Fail at rank 1: KB-04 first, KB-03 second'],['4','Weather in Pune tomorrow','Abstain without evidence','Pass in recorded offline baseline'],['5','New Bluetooth / charging system with Qwen and PostgreSQL','Generate a new structured draft and persist it','Not run end to end; local execution pending']], [.3,2,2,2.3])
table(D,[['Metric','Score / observation'],['Offline top-1 expected-source accuracy','87.5% (14/16 positive questions)'],['Offline retrieval recall at 3','93.75% (15/16)'],['Offline MRR at 3','0.90625'],['Offline paraphrase top-1','50% (2/4)'],['Offline out-of-domain abstention','2/2 questions'],['Mean offline evaluation latency','0.845 ms in recorded small-corpus run; not LLM latency'],['Software tests before final update','38 passed in prior run; AI adapter tests used simulated model responses. These are not real-model quality results.'],['LLM faithfulness and PostgreSQL runtime','Not measured. No new result or accuracy claim is supplied.']], [2.25,4.35])
para(D,'Source evidence: Evaluation_Results/metrics.json, retrieval_results.json, test_run.txt and test_run_v1_1.txt. Live generation output should be exported from the AI Assessment page after local setup. The optional verification utility is supplied but was not run for this final update.')
page(D);h(D,'6. Responsible AI Measures')
table(D,[['Area','Measure adopted in this project'],['Grounding and citations','Snapshot submitted input and retrieved passages; require known source IDs and input citations. Distinguish source existence from factual support.'],['Privacy','Use synthetic or approved inputs. Ollama is local. No external cloud LLM API is required. Model and container downloads occur during setup.'],['Security','Validate structured output; parameterize SQL; escape rendered text; reject cross-origin writes; keep services local. No production authentication or tamper-proof audit is claimed.'],['Human oversight / review','All LLM assessments are drafts. Unknown controls are not treated as absent. Human reviewers evaluate assumptions, controls and proposed requirements; risk stays Unassessed.']], [1.6,5])
h(D,'7. Innovation Highlights')
para(D,'The main extension is a separate generative assessment workflow for new system descriptions. Unlike the baseline tag-to-pattern demonstration, this route asks the model to identify assets and propose context-specific scenarios. It retains an evidence snapshot, model digest and review state for each run.')
para(D,'The application combines a reviewable AI output with structured persistence and explicit failure reporting. A new assessment can be reopened, reviewed and exported without regenerating it. The existing version-comparison workflow demonstrates how changes can invalidate older baseline findings. Automatic change propagation from free-text AI assessments into that graph is not implemented.')
para(D,'The approach is intentionally compact for a student implementation: one service layer, a small local model, a single PostgreSQL service for records and vectors, and a browser interface. Its value depends on meaningful source data and reviewer judgment; it is not evidence of superior threat coverage.')
page(D);h(D,'8. Screenshots, Logs and Evidence')
para(D,'The following screenshots were captured from the earlier offline application. They demonstrate retrieval and change analysis only, not live Qwen generation. New AI-page output and PostgreSQL evidence must be captured after running the updated package locally.')
pic(D,R/'Evaluation_Results/screenshots/08_grounded_answer.png',5.3);para(D,'Figure 1  Offline evidence explorer returning the KB-02 rollback passage with its source reference.')
pic(D,R/'Evaluation_Results/screenshots/11_changes.png',5.3);para(D,'Figure 2  Baseline architecture comparison showing changed entities and affected prior findings.')
page(D);h(D,'9. Reflection and Conclusions')
para(D,'The baseline established data handling, traceability and review behavior but relied on prepared threat content. Version 1.1 adds the missing generative route: a fresh system description is sent to a local instruction model, and structured draft output is saved for review. This improves the scope of the implementation without turning generated advice into an approved assessment.')
para(D,'The main remaining uncertainty is real-model quality and deployment behavior on the evaluation machine. A small model can omit threats, infer unsupported controls or produce an invalid schema. Lexical retrieval also struggles with paraphrases, as the recorded failures show. Future improvements include domain-reviewed examples, learned semantic embeddings and expert-scored independent evaluations.')
h(D,'10. AI-Tool Usage Declaration')
table(D,[['AI tool / model','Purpose','Artifact affected','Verified By Student Y/N'],['ChatGPT / Codex','Coding assistance, synthetic data, documentation and configuration drafting','Source, test scripts, input examples, report and synopsis','Pending student review'],['Qwen2.5 3B Instruct via Ollama','Runtime structured automotive assessment generation','Generated assets, threats, mitigations and requirements','Not run in this final update']], [1.5,1.8,2,1.3])
para(D,'Students must record their actual implementation, review and demonstration contributions. Generated assistance must not be presented as independently authored or empirically validated work without supporting evidence.')
h(D,'11. Student Declaration')
para(D,'I confirm that this report represents my own understanding and contribution, that the work follows the approved case-study scope, and that I can explain every significant part of the implementation.')
table(D,[['Signature:','________________','Date:','________________'],['Name:','________________','PRN:','________________']], [1,2.3,1,2.3])
para(D,'Complete and sign after reviewing the implementation and obtaining the required faculty approval.')
para(D,'Technical references: FastAPI testing documentation (fastapi.tiangolo.com/tutorial/testing/); Ollama structured outputs and model listing (docs.ollama.com, ollama.com/library/qwen2.5:3b-instruct); pgvector official repository (github.com/pgvector/pgvector); supplied Automotive Engineering AI case study document, CS3. Accessed 9 October 2026.')
D.save(R/'Documentation/Technical_Report.docx')
S=new(str(R/'Documentation/templates/Synopsis_Template.docx'));brand(S,'PROJECT SYNOPSIS')
h(S,'Student Details')
for t in ['Institute/University: PCCOE','Programme / Batch: ______________________________','Student Name: ______________________________','PRN: ______________________________']:para(S,t)
h(S,'1. Case Study ID and Project Title / Topic');para(S,'Case Study ID: CS3');para(S,'Project Title / Topic: TARA Studio Automotive Cybersecurity Assessment Assistant')
h(S,'2. Objectives')
for t in ['Extract candidate cybersecurity assets from a new automotive system description using a local instruction model.','Generate evidence-linked draft threat and damage scenarios with assumptions, mitigations, requirements and verification proposals.','Persist model outputs and review history in PostgreSQL and make each assessment retrievable and exportable.','Maintain clear limits through citations, schema checks, unassessed risk and human review.']:S.add_paragraph(t,style='List Bullet')
h(S,'3. Scope');para(S,'In Scope – To Be Implemented',True);para(S,'A browser-based AI Assessment page; local Ollama inference; retrieval from synthetic or approved reference content; structured output checks; PostgreSQL persistence; review and JSON export. The offline baseline also supplies architecture version comparison and educational risk support.')
para(S,'Out of Scope – Will Not Be Implemented',True);para(S,'Autonomous cybersecurity acceptance, certified standards scoring, live vehicle exploitation, arbitrary diagram understanding, production multi-user security and automatic conversion of generated text into an approved architecture graph.')
page(S);h(S,'4. Proposed Approach')
for title,body in [('Pre-trained model name version and hosting','Qwen2.5 3B Instruct, Ollama tag qwen2.5:3b-instruct, local inference. Capture the installed digest; weights are downloaded separately.'),('Embedding model','No pretrained embedding model. The initial retriever uses BM25 with deterministic 384-dimensional token-hash vectors.'),('Vector store FAISS or ChromaDB','PostgreSQL with pgvector is selected instead of FAISS/ChromaDB; SQLite is a labelled local fallback.'),('RAG pipeline and AI-agent design','Snapshot the description as SYS-xx sources, retrieve up to four relevant reference passages, generate a structured assessment and validate identifiers before saving. No autonomous tool execution.'),('Service layer and UI','FastAPI and Python backend, Uvicorn server, Pydantic schemas and local HTML/CSS/JavaScript interface.'),('Tools and deterministic rules','pypdf ingestion, parameterized PostgreSQL writes, schema and citation checks, job states, review revision control, baseline educational risk matrix.'),('Proposed System / Architecture','The application, Ollama and PostgreSQL run as separate local services using Docker Compose.')]:
 para(S,title,True);para(S,body)
pic(S,R/'Documentation/figures/architecture.png',6.6)
page(S);h(S,'5. Input Data / Knowledge Base')
para(S,'Data / Knowledge Sources',True);para(S,'Original synthetic OTA architectures and a 12-passage reference handbook, plus newly entered system descriptions. Bluetooth vehicle access and EV charging descriptions are supplied as additional inputs. Use approved or synthetic data only.')
para(S,'Approximate Size and Format',True);para(S,'Two architecture JSON versions, five baseline threat patterns, twelve knowledge passages, eighteen baseline questions, and two text descriptions. AI input is limited to 8,000 characters per assessment.')
para(S,'Preprocessing and EDA Planned',True);para(S,'Extract text and preserve source metadata. Use 1,200-character knowledge windows with 100-character overlap; create 1,000-character non-overlapping submitted-input chunks. Validate references and distinguish assumptions from facts.')
h(S,'6. Expected Outcomes and Evaluation Plan')
para(S,'Expected Outputs',True);para(S,'A saved draft assessment containing extracted assets, threats, damage scenarios, hypotheses, controls, assumptions, evidence IDs, mitigations, requirements and verification proposals. Risk remains Unassessed until a suitable human-led method is applied.')
para(S,'Evaluation Method',True);para(S,'Use the recorded synthetic retrieval benchmark as a baseline. For a later live-model evaluation, inspect generated results from new system descriptions, source support, omitted threats, valid identifiers, persistence and latency. Do not equate software fixture tests with model quality; no further validations are claimed in this final update.')
para(S,'Responsible AI Measures',True);para(S,'Local processing, approved or synthetic inputs, source snapshots, explicit failure states, no fabricated citations, no autonomous approval and documented human review. Factual support still requires review.')
para(S,'MLOps / Deployment Elements Planned',True);para(S,'Versioned prompts and configuration, containerized services, persistent database/model volumes, model digest capture, structured job records and JSON export.')
h(S,'7. Planned AI-Tool Usage')
para(S,'AI Tool / Model: ChatGPT / Codex for development assistance; Qwen2.5 3B Instruct via Ollama for runtime generation.')
para(S,'Purpose: Assist source-code and document preparation, and generate structured candidate automotive assessments.')
para(S,'Artifact Affected: Application code, synthetic inputs, prompts, report, synopsis and generated assessment records. Students must disclose their actual contributions and review.')
page(S);h(S,'Student Declaration')
para(S,'[ ] I will not change the approved case-study scope, major methodology, or implementation approach without prior faculty approval.')
para(S,'[ ] I confirm that the work submitted under this synopsis represents my own understanding and contribution, and I will be responsible for explaining every significant part of the selected case study and its implementation.')
para(S,'Name and Signature: ______________________________    Date: ________________')
h(S,'Faculty Approval Record')
para(S,'Faculty guide: __________________________________________')
para(S,'Decision: [ ] Approved   [ ] Approved with changes   [ ] Revise')
para(S,'Comments: ______________________________________________')
para(S,'________________________________________________________')
para(S,'Faculty signature: ________________________ Date: ___________')
para(S,'Approval is pending until the faculty member reviews and signs this record.')
S.save(R/'Synopsis/Project_Synopsis.docx')
print('Filled report and synopsis templates')
